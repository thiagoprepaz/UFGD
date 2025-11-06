#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Automação Progressão por Mérito — GUI (.ods -> resultado_progressao.xlsx)

Regras por planilha:
* Se J13 != 0 -> linha: A13, B13, MES/ANO=C5, rubrica="00001", rendimento="r",
  sequência=C6, valor=abs(J13), justificativa=C9, documento legal=C10
* Se N13 != 0 -> idem, mas com valor=abs(N13)

Saída:
- resultado_progressao.xlsx na mesma pasta do .ods

Dependências:
    pip install pandas odfpy openpyxl
"""

import os
import re
import sys
import platform
from typing import Optional, Any, List, Dict

# ---- imports com mensagens amigáveis ----
try:
    import pandas as pd
except ModuleNotFoundError:
    import tkinter as _tk
    from tkinter import messagebox as _mb
    _tk.Tk().withdraw()
    _mb.showerror(
        "Dependência ausente",
        "O módulo 'pandas' não está instalado.\n\nInstale com:\n\npip install pandas odfpy openpyxl"
    )
    sys.exit(1)

# Tkinter
import tkinter as tk
from tkinter import filedialog, messagebox

# ---- Utilidades de célula/parse ----
def col_to_index(col: str) -> int:
    col = col.strip().upper()
    idx = 0
    for ch in col:
        if not ('A' <= ch <= 'Z'):
            raise ValueError(f"Coluna inválida: {col}")
        idx = idx * 26 + (ord(ch) - ord('A') + 1)
    return idx - 1

def get_cell(df: pd.DataFrame, col_letter: str, row_number: int) -> Optional[Any]:
    r = row_number - 1
    c = col_to_index(col_letter)
    try:
        return df.iat[r, c]
    except Exception:
        return None

def parse_br_number(value: Any) -> Optional[float]:
    """
    Converte 'R$ 1.234,56' (ou similares) para float.
    Retorna None se vazio/inválido.
    """
    if value is None:
        return None
    if isinstance(value, (int, float)):
        try:
            return float(value)
        except Exception:
            return None
    s = str(value).strip()
    if s == "" or s.lower() in {"nan", "none"}:
        return None
    s = s.replace("R$", "").replace("r$", "").replace(" ", "").replace("\xa0", "")
    s = s.replace("%", "")
    s = s.replace(".", "")
    s = s.replace(",", ".")
    s = re.sub(r"[^0-9\.\-]", "", s)
    try:
        return float(s)
    except Exception:
        return None

# ---- Núcleo de processamento ----
def process_sheet(df: pd.DataFrame) -> List[Dict[str, Any]]:
    linhas: List[Dict[str, Any]] = []

    a13 = get_cell(df, "A", 13)
    b13 = get_cell(df, "B", 13)
    mes_ano_c5 = get_cell(df, "C", 5)   # MES/ANO
    seq_c6 = get_cell(df, "C", 6)
    justificativa_c9 = get_cell(df, "C", 9)
    doc_legal_c10 = get_cell(df, "C", 10)

    j13 = parse_br_number(get_cell(df, "J", 13))
    if j13 is not None and j13 != 0:
        linhas.append({
            "A13": a13,
            "B13": b13,
            "MES/ANO": mes_ano_c5,      # entre B13 e rubrica
            "rubrica": "00001",
            "rendimento": "r",
            "sequência": seq_c6,
            "valor": abs(float(j13)),    # sempre positivo
            "justificativa": justificativa_c9,
            "documento legal": doc_legal_c10,
        })

    n13 = parse_br_number(get_cell(df, "N", 13))
    if n13 is not None and n13 != 0:
        linhas.append({
            "A13": a13,
            "B13": b13,
            "MES/ANO": mes_ano_c5,
            "rubrica": "00001",
            "rendimento": "r",
            "sequência": seq_c6,
            "valor": abs(float(n13)),    # sempre positivo
            "justificativa": justificativa_c9,
            "documento legal": doc_legal_c10,
        })

    return linhas

def build_table_from_ods(file_path: str) -> pd.DataFrame:
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Arquivo não encontrado: {file_path}")

    # Descobrir planilhas via engine 'odf'
    try:
        xls = pd.ExcelFile(file_path, engine="odf")
    except Exception as e:
        raise RuntimeError(
            "Falha ao abrir o .ods. Verifique se 'odfpy' está instalado:\n"
            "pip install odfpy\n\n"
            f"Detalhe: {e}"
        ) from e

    todas_linhas: List[Dict[str, Any]] = []
    for nome in xls.sheet_names:
        # header=None para ler grade “crua” (posições exatas)
        df = pd.read_excel(file_path, sheet_name=nome, header=None, engine="odf")
        todas_linhas.extend(process_sheet(df))

    colunas = [
        "A13", "B13", "MES/ANO", "rubrica",
        "rendimento", "sequência", "valor",
        "justificativa", "documento legal"
    ]
    return pd.DataFrame(todas_linhas, columns=colunas)

def salvar_excel(df: pd.DataFrame, origem: str) -> str:
    """
    Salva apenas XLSX (sem CSV) na mesma pasta do .ods.
    Formata a coluna 'valor' como moeda (se possível).
    """
    # Garante 'valor' numérico e em módulo (fallback de segurança)
    if "valor" in df.columns:
        df["valor"] = pd.to_numeric(df["valor"], errors="coerce").abs()

    base_dir = os.path.dirname(os.path.abspath(origem))
    xlsx_out = os.path.join(base_dir, "resultado_progressao.xlsx")

    try:
        with pd.ExcelWriter(xlsx_out, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="resultado")

            # Formatação opcional da coluna 'valor' como moeda BRL
            try:
                from openpyxl.utils import get_column_letter
                ws = writer.sheets["resultado"]
                if "valor" in df.columns:
                    col_idx = df.columns.get_loc("valor") + 1  # 1-based
                    numero_col = get_column_letter(col_idx)
                    for row in range(2, len(df) + 2):
                        ws[f"{numero_col}{row}"].number_format = 'R$ #,##0.00'
                # Ajuste simples de largura de coluna
                for i, col_name in enumerate(df.columns, start=1):
                    maxlen = max(
                        (len(str(x)) for x in [col_name] + df[col_name].astype(str).tolist()),
                        default=10
                    )
                    ws.column_dimensions[get_column_letter(i)].width = min(max(10, maxlen + 2), 60)
            except Exception:
                # Se algo falhar na formatação, ignorar sem travar
                pass

    except Exception as e:
        raise RuntimeError(
            f"Falha ao salvar XLSX ({e}). Verifique se 'openpyxl' está instalado:\n"
            "pip install openpyxl"
        ) from e

    return xlsx_out

def abrir_pasta(path: str):
    try:
        if platform.system() == "Windows":
            os.startfile(path)  # type: ignore[attr-defined]
        elif platform.system() == "Darwin":
            import subprocess
            subprocess.Popen(["open", path])
        else:
            import subprocess
            subprocess.Popen(["xdg-open", path])
    except Exception:
        pass

# ---- Interface gráfica ----
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Automação Progressão por Mérito")
        self.geometry("560x260")
        self.resizable(False, False)

        self.ods_path: Optional[str] = None

        self.lbl_info = tk.Label(self, text="Selecione um arquivo .ods para processar:", anchor="w")
        self.lbl_info.pack(fill="x", padx=12, pady=(12, 6))

        frm = tk.Frame(self)
        frm.pack(fill="x", padx=12)

        self.ent_path = tk.Entry(frm)
        self.ent_path.pack(side="left", fill="x", expand=True, padx=(0, 8))

        self.btn_browse = tk.Button(frm, text="Escolher .ods", command=self.escolher_arquivo)
        self.btn_browse.pack(side="left")

        self.btn_process = tk.Button(self, text="Processar", command=self.processar, state="disabled", height=2)
        self.btn_process.pack(fill="x", padx=12, pady=10)

        self.lbl_status = tk.Label(self, text="", anchor="w", fg="#333")
        self.lbl_status.pack(fill="x", padx=12)

        self.btn_abrir_pasta = tk.Button(self, text="Abrir pasta de saída", command=self.abrir_saida, state="disabled")
        self.btn_abrir_pasta.pack(padx=12, pady=10)

        self.lbl_rodape = tk.Label(self, text="Versão 1.3 - GUI (somente XLSX)", anchor="e", fg="#777")
        self.lbl_rodape.pack(fill="x", padx=12, pady=(8, 10))

    def escolher_arquivo(self):
        caminho = filedialog.askopenfilename(
            title="Selecione o arquivo .ods",
            filetypes=[("Planilhas ODS", "*.ods"), ("Todos os arquivos", "*.*")]
        )
        if caminho:
            self.ods_path = caminho
            self.ent_path.delete(0, tk.END)
            self.ent_path.insert(0, caminho)
            self.btn_process.config(state="normal")
            self.lbl_status.config(text="Arquivo selecionado. Pronto para processar.")

    def processar(self):
        if not self.ods_path:
            messagebox.showwarning("Aviso", "Escolha um arquivo .ods primeiro.")
            return

        try:
            self.lbl_status.config(text="Processando, aguarde...")
            self.update_idletasks()

            df = build_table_from_ods(self.ods_path)
            xlsx_path = salvar_excel(df, self.ods_path)

            linhas = len(df)
            msg = "Processamento concluído.\n"
            msg += f"Linhas geradas: {linhas}\n"
            msg += f"XLSX: {xlsx_path}\n"

            self.lbl_status.config(text=msg)
            self.btn_abrir_pasta.config(state="normal")
            messagebox.showinfo("Concluído", msg)

        except Exception as e:
            messagebox.showerror(
                "Erro ao processar",
                f"Ocorreu um erro:\n\n{e}\n\nDica: garanta que as dependências estejam instaladas:\n"
                "pip install pandas odfpy openpyxl"
            )
            self.lbl_status.config(text=f"Erro: {e}")

    def abrir_saida(self):
        if self.ods_path:
            abrir_pasta(os.path.dirname(os.path.abspath(self.ods_path)))

if __name__ == "__main__":
    app = App()
    app.mainloop()
